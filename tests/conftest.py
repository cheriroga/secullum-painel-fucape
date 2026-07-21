import openpyxl
import pytest


@pytest.fixture(autouse=True)
def _isolar_variaveis_de_ambiente_do_app(monkeypatch):
    """Garante que nenhum teste dependa do .env real do desenvolvedor.

    webapp.main chama load_dotenv() na importação, o que deixa essas
    variáveis no os.environ pelo resto da sessão de testes — sem isso,
    o valor real de PAINEL_METODO_ENVIO (ou qualquer outra) no .env de
    quem está rodando os testes vaza pros testes que não a definem
    explicitamente."""
    for nome in (
        "PAINEL_METODO_ENVIO", "PAINEL_MODO_TESTE",
        "GRAPH_TENANT_ID", "GRAPH_CLIENT_ID", "GRAPH_CLIENT_SECRET", "GRAPH_REMETENTE",
    ):
        monkeypatch.delenv(nome, raising=False)


@pytest.fixture
def workbook_path(tmp_path):
    def _construir(blocos):
        """`blocos` é uma lista de dicts com chaves:
        nome, funcao, admissao, departamento, dias.
        `dias` é uma lista de tuplas de 12 elementos:
        (data_str "DD/MM/AAAA", sufixo_dia_semana, ent1, sai1, ent2, sai2, ent3, sai3,
         ex50, atraso, btotal, exnot).
        Batidas (ent/sai) são datetime.timedelta (hora real), str (código de status,
        ex: "FALTA") ou None. ex50/atraso/exnot são datetime.timedelta ou None.
        btotal é string "+HH:MM"/"-HH:MM" ou None — nunca int/float cru (o xlsx real
        nunca guarda isso)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        linha = 1
        for bloco in blocos:
            ws.cell(row=linha, column=1, value="CARTÃO PONTO")
            ws.cell(row=linha + 1, column=1, value="Período: 01/01/2026 até 09/07/2026.")
            ws.cell(row=linha + 3, column=1, value="Empresa")
            ws.cell(row=linha + 3, column=2, value="FUCAPE PESQUISA E ENSINO SA")
            ws.cell(row=linha + 4, column=1, value="CNPJ")
            ws.cell(row=linha + 5, column=1, value="Inscrição")
            ws.cell(row=linha + 5, column=2, value="ISENTO")
            ws.cell(row=linha + 7, column=1, value="Nome")
            ws.cell(row=linha + 7, column=2, value=bloco["nome"])
            ws.cell(row=linha + 8, column=1, value="Nº Identificador")
            ws.cell(row=linha + 9, column=1, value="C.T.P.S.")
            ws.cell(row=linha + 9, column=4, value="Admissão")
            ws.cell(row=linha + 9, column=5, value=bloco["admissao"] or "")
            ws.cell(row=linha + 10, column=1, value="Função")
            ws.cell(row=linha + 10, column=2, value=bloco["funcao"])
            ws.cell(row=linha + 11, column=1, value="Departamento")
            ws.cell(row=linha + 11, column=2, value=bloco["departamento"] or "")
            ws.cell(row=linha + 12, column=1, value="Observação")

            header_row = linha + 15  # Nome (linha+7) + offset 8, igual ao arquivo real
            ws.cell(row=header_row, column=1, value="Data")
            totais_row = header_row + 1
            ws.cell(row=totais_row, column=1, value="Totais")

            r = totais_row + 1
            for dia in bloco["dias"]:
                data_str, sufixo, ent1, sai1, ent2, sai2, ent3, sai3, ex50, atraso, btotal, exnot = dia
                ws.cell(row=r, column=1, value=f"{data_str} - {sufixo}")
                ws.cell(row=r, column=2, value=ent1)
                ws.cell(row=r, column=3, value=sai1)
                ws.cell(row=r, column=4, value=ent2)
                ws.cell(row=r, column=5, value=sai2)
                ws.cell(row=r, column=6, value=ent3)
                ws.cell(row=r, column=7, value=sai3)
                ws.cell(row=r, column=8, value=ex50)
                ws.cell(row=r, column=9, value=atraso)
                ws.cell(row=r, column=10, value=btotal)
                ws.cell(row=r, column=11, value=exnot)
                r += 1

            linha = r + 10  # espaço em branco antes do próximo bloco (legenda/assinatura no arquivo real)
        caminho = tmp_path / "cartaoponto_teste.xlsx"
        wb.save(caminho)
        return caminho
    return _construir
