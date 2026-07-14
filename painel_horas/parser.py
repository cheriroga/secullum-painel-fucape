import re
from dataclasses import dataclass, field
from datetime import date

import openpyxl

from painel_horas.horas import parse_horas

_PADRAO_DIA = re.compile(r"^(\d{2}/\d{2}/\d{4}) - (.+)$")


@dataclass
class Dia:
    data: date
    dia_semana: str
    ent1: object
    sai1: object
    ent2: object
    sai2: object
    ent3: object
    sai3: object
    ex50_min: int
    atraso_min: int
    btotal_min: int | None
    exnot_min: int


@dataclass
class Colaborador:
    nome: str
    funcao: str
    admissao: date | None
    departamento: str
    dias: list[Dia] = field(default_factory=list)
    total_bruto_min: int = 0


def _parse_data(texto: str) -> date:
    dia, mes, ano = texto.strip().split("/")
    return date(int(ano), int(mes), int(dia))


def ler_colaboradores(caminho_xlsx) -> list[Colaborador]:
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    max_row = ws.max_row

    def val(linha, coluna):
        return ws.cell(row=linha, column=coluna).value

    colaboradores: list[Colaborador] = []
    r = 1
    while r <= max_row:
        if val(r, 1) != "Nome":
            r += 1
            continue

        nome_bruto = val(r, 2)
        nome = nome_bruto.strip() if isinstance(nome_bruto, str) else ""
        if not nome:
            print(f"[aviso] bloco na linha {r} sem nome reconhecível — pulando")
            r += 1
            continue

        admissao_str = val(r + 2, 5)
        admissao = (
            _parse_data(admissao_str)
            if isinstance(admissao_str, str) and admissao_str.strip()
            else None
        )
        funcao_bruta = val(r + 3, 2)
        funcao = funcao_bruta.strip() if isinstance(funcao_bruta, str) else ""
        departamento_bruto = val(r + 4, 2)
        departamento = (
            departamento_bruto.strip()
            if isinstance(departamento_bruto, str) and departamento_bruto.strip()
            else "Sem Departamento"
        )
        if admissao is None or departamento == "Sem Departamento":
            print(f"[aviso] colaborador '{nome}' com admissão/departamento incompleto — usando fallback")

        header_row = None
        limite = min(max_row, r + 15)
        j = r + 5
        while j <= limite:
            if val(j, 1) == "Data":
                header_row = j
                break
            j += 1

        dias: list[Dia] = []
        if header_row is None:
            print(f"[aviso] colaborador '{nome}' sem tabela diária reconhecível — pulando dias")
            k = r + 1
        else:
            k = header_row + 2  # pula a linha "Totais"
            while k <= max_row:
                bruto = val(k, 1)
                m = _PADRAO_DIA.match(bruto) if isinstance(bruto, str) else None
                if not m:
                    break
                btotal_raw = val(k, 10)
                dias.append(Dia(
                    data=_parse_data(m.group(1)),
                    dia_semana=m.group(2),
                    ent1=val(k, 2), sai1=val(k, 3),
                    ent2=val(k, 4), sai2=val(k, 5),
                    ent3=val(k, 6), sai3=val(k, 7),
                    ex50_min=parse_horas(val(k, 8)),
                    atraso_min=parse_horas(val(k, 9)),
                    btotal_min=parse_horas(btotal_raw) if btotal_raw is not None else None,
                    exnot_min=parse_horas(val(k, 11)),
                ))
                k += 1

        total_bruto = sum(d.btotal_min or 0 for d in dias)
        colaboradores.append(Colaborador(
            nome=nome, funcao=funcao, admissao=admissao, departamento=departamento,
            dias=dias, total_bruto_min=total_bruto,
        ))
        r = k

    return colaboradores
